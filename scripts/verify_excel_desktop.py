"""Controlled Windows Excel checks on generated files only; never attaches to user's workbook.

Optional dependency: pywin32. DispatchEx owns a separate hidden Excel instance.
The main process places a time limit around the worker and only targets its own Excel PID.
"""

import argparse
import json
import os
import subprocess
import sys
import time
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook

from rf_link_calculator.domain.models import NoiseSpec, PowerSpec, TwoTone
from rf_link_calculator.engine import calculate
from rf_link_calculator.examples import amp, examples
from rf_link_calculator.exporters.excel import build_excel
from rf_link_calculator.persistence.json_io import atomic_write


def scenarios():
    p = examples()["两级放大器"]
    cases = [("baseline", p, [])]
    for name, field, value, cell in [
        ("pin", "input_power_dbm", -20, "C2"),
        ("bandwidth", "noise_bandwidth_hz", 1e7, "C5"),
        ("temperature", "source_noise_temperature_k", 100, "C6"),
    ]:
        cases.append(
            (name, replace(p, analysis=replace(p.analysis, **{field: value})), [("输入配置", cell, value)])
        )
    for name, changes, col, value in [
        ("gain", {"gain_db": 15}, "F4", 15),
        ("nf", {"noise": NoiseSpec("manual", 4)}, "H4", 4),
        ("ip3", {"ip3": PowerSpec("finite", "input", 15)}, "L4", 15),
        ("p1", {"p1db": PowerSpec("finite", "input", -5)}, "O4", -5),
        ("missing_nf", {"noise": NoiseSpec("manual", None)}, "H4", None),
        ("missing_p1", {"p1db": PowerSpec("finite", "input", None)}, "O4", None),
        ("disabled", {"enabled": False}, "B4", False),
        ("mixed_p", {"compression_p": 7}, "P4", 7),
    ]:
        cases.append(
            (
                name,
                replace(p, stages=(replace(p.stages[0], **changes), p.stages[1])),
                [("器件参数", col, value)],
            )
        )
    reordered = replace(p, stages=(replace(p.stages[1], order=1), replace(p.stages[0], order=2)))
    cases.append(("reorder", reordered, [("swap", "A4:Y4", "A5:Y5")]))
    appended = replace(p, stages=(*p.stages, amp("新级", 5, 2, 10, 0, 3)))
    cases.append(("add", appended, [("copy_row", 3, appended)]))
    deleted = replace(p, stages=(p.stages[1],))
    cases.append(("delete_by_clear", deleted, [("clear_row", "A4:Y4", None)]))
    missing_p1 = replace(p, stages=(replace(p.stages[0], p1db=PowerSpec("unknown")), p.stages[1]))
    cases.append(("invalid_default_p", missing_p1, [("器件参数", "P4", None), ("输入配置", "C11", 0)]))
    cases.append(
        (
            "two_tone_high",
            replace(p, analysis=replace(p.analysis, two_tone=TwoTone(True, 10, 1e5))),
            [("输入配置", "C13", 10)],
        )
    )
    for name, key in [
        ("receiver_B", "接收链路"),
        ("compression_C", "累计压缩反例"),
        ("passive_D", "无源衰减器"),
        ("temperature_E", "源温度600K"),
        ("mixer_F", "下变频链路"),
        ("ideal", "全部理想线性"),
        ("all_disabled", "全部禁用"),
        ("mixer_noise_unknown", "混频噪声口径未知"),
        ("long_200", "200级长链路"),
    ]:
        cases.append((name, examples()[key], None))
    return p, cases


def worker(output: Path):
    import pythoncom
    import win32api
    import win32com.client
    import win32process

    pythoncom.CoInitialize()
    excel, book = None, None
    report = {"status": "running", "scenarios": [], "started_at": time.time()}
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AskToUpdateLinks = False
        excel.AutomationSecurity = 3
        pid = win32process.GetWindowThreadProcessId(excel.Hwnd)[1]
        handle = win32api.OpenProcess(0x0410, False, pid)
        executable = win32process.GetModuleFileNameEx(handle, 0)
        created = str(win32process.GetProcessTimes(handle)["CreationTime"])
        handle.Close()
        atomic_write(
            output / "office-owned-process.json",
            json.dumps({"pid": pid, "executable": executable, "created": created}).encode("utf-8"),
        )
        report["excel_version"] = excel.Version
        report["excel_build"] = str(excel.Build)
        report["automation_reported_name"] = excel.Name
        report["application_path"] = excel.Path
        report["executable"] = executable
        report["actual_engine"] = (
            "WPS Spreadsheet"
            if Path(executable).name.lower() == "et.exe"
            else ("Microsoft Excel" if Path(executable).name.lower() == "excel.exe" else "unverified engine")
        )
        p, cases = scenarios()
        base = build_excel(p, calculate(p))
        for name, changed, edits in cases:
            file = output / (name + ".xlsx")
            atomic_write(file, base if edits is not None else build_excel(changed, calculate(changed)))
            book = excel.Workbooks.Open(str(file.resolve()), UpdateLinks=0, ReadOnly=False)
            for sheet, cell, value in edits or []:
                if sheet == "swap":
                    ws = book.Worksheets("器件参数")
                    first, second = ws.Range(cell).Value, ws.Range(value).Value
                    ws.Range(cell).Value, ws.Range(value).Value = second, first
                elif sheet == "clear_row":
                    book.Worksheets("器件参数").Range(cell).ClearContents()
                elif sheet == "copy_row":
                    s = value.stages[cell - 1]
                    vals = (
                        cell,
                        s.enabled,
                        s.id,
                        s.name,
                        s.type,
                        s.gain_db,
                        s.noise.mode,
                        s.noise.value_db,
                        s.physical_temperature_k,
                        s.ip3.mode,
                        s.ip3.reference,
                        s.ip3.value_dbm,
                        s.p1db.mode,
                        s.p1db.reference,
                        s.p1db.value_dbm,
                        s.compression_p,
                        "构造新增级",
                        None,
                        None,
                        None,
                        "",
                        "unknown",
                        False,
                        None,
                        "",
                    )
                    book.Worksheets("器件参数").Range(f"A{cell + 3}:Y{cell + 3}").Value = (vals,)
                else:
                    target = book.Worksheets(sheet).Range(cell)
                    if value is None:
                        target.ClearContents()
                    else:
                        target.Value = value
            excel.CalculateFullRebuild()
            book.Save()
            book.Close(SaveChanges=False)
            book = None
            wb = load_workbook(file, data_only=True)
            expected = calculate(changed)
            mismatches = []
            for row in wb["系统汇总"].iter_rows(min_row=4, max_row=33, values_only=True):
                _, actual, _, state, key = row[:5]
                if key not in expected.metrics:
                    continue
                metric = expected.metrics[key]
                if metric.value is None:
                    if isinstance(actual, (int, float)):
                        mismatches.append({"key": key, "actual": actual, "expected": None, "state": state})
                elif not isinstance(actual, (float, int)) or abs(actual - metric.value) > 1e-5:
                    mismatches.append(
                        {"key": key, "actual": actual, "expected": metric.value, "state": state}
                    )
            errors = []
            for ws in wb:
                for row in ws:
                    for c in row:
                        if c.data_type == "e":
                            errors.append(f"{ws.title}!{c.coordinate}: {c.value}")
            result = {
                "name": name,
                "pass": not mismatches and not errors,
                "mismatches": mismatches,
                "formula_errors": errors[:30],
            }
            report["scenarios"].append(result)
            print(json.dumps(result, ensure_ascii=False), flush=True)
            atomic_write(
                output / "excel-verification.json",
                json.dumps(report, ensure_ascii=False, indent=2).encode("utf-8"),
            )
        report["status"] = "passed" if all(s["pass"] for s in report["scenarios"]) else "failed"
    except Exception as exc:
        report["status"] = "environment_or_execution_failure"
        report["error"] = f"{type(exc).__name__}: {exc}"
        print(report["error"], flush=True)
    finally:
        if book is not None:
            book.Close(SaveChanges=False)
        if excel is not None:
            excel.Quit()
        pythoncom.CoUninitialize()
        atomic_write(
            output / "excel-verification.json",
            json.dumps(report, ensure_ascii=False, indent=2).encode("utf-8"),
        )
    return 0 if report["status"] == "passed" else 1


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path("artifacts/excel-desktop"))
    parser.add_argument("--worker", action="store_true")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    if args.worker:
        return worker(args.output.resolve())
    env = dict(os.environ, PYTHONUTF8="1")
    try:
        return subprocess.run(
            [sys.executable, __file__, "--worker", "--output", str(args.output.resolve())],
            env=env,
            timeout=600,
        ).returncode
    except subprocess.TimeoutExpired:
        # Only terminate the Excel PID explicitly recorded by our DispatchEx instance.
        pidfile = args.output / "office-owned-process.json"
        if pidfile.exists():
            import pywintypes
            import win32api
            import win32process

            owned = json.loads(pidfile.read_text("utf-8"))
            try:
                handle = win32api.OpenProcess(0x0411, False, owned["pid"])
                if (
                    win32process.GetModuleFileNameEx(handle, 0) == owned["executable"]
                    and str(win32process.GetProcessTimes(handle)["CreationTime"]) == owned["created"]
                ):
                    win32api.TerminateProcess(handle, 1)
                handle.Close()
            except pywintypes.error:
                pass  # Already exited; never target a different process.
        print("Excel验证超时，已清理本次独立实例；未触及用户已有Excel进程。")
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
