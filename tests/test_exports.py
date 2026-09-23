import json
from dataclasses import replace
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from rf_link_calculator.application.project_service import edit_structure
from rf_link_calculator.diagram.scene import build_scene, segmented_scenes
from rf_link_calculator.engine import calculate
from rf_link_calculator.examples import examples
from rf_link_calculator.exporters.bundle import export_bundle, safe_name
from rf_link_calculator.exporters.drawio import render_drawio
from rf_link_calculator.exporters.excel import SHEETS, build_excel
from rf_link_calculator.exporters.svg import render_svg
from rf_link_calculator.persistence.csv_io import export_csv, import_csv
from rf_link_calculator.persistence.json_io import dumps, loads, project_from_dict, save_project


def test_json_roundtrip_and_bounds(tmp_path):
    for p in examples().values():
        assert loads(dumps(p)) == p
    p = examples()["接收链路"]
    for val in (float("nan"), float("inf"), True, "NaN"):
        raw = p.to_dict()
        raw["analysis"]["input_power_dbm"] = val
        with pytest.raises(ValueError):
            project_from_dict(raw)
    with pytest.raises(ValueError):
        loads('{"schema_version":"2.0.0"}')
    with pytest.raises(ValueError):
        loads('{"notes":"a","notes":"b"}')
    with pytest.raises(ValueError):
        loads(b" " * (5 * 1024 * 1024 + 1))
    d = p.to_dict()
    d["stages"][1]["id"] = d["stages"][0]["id"]
    with pytest.raises(ValueError):
        project_from_dict(d)
    target = tmp_path / "中文.json"
    save_project(p, target)
    changed = replace(p, notes="changed")
    save_project(changed, target)
    assert loads(target.read_bytes()) == changed
    assert loads(next(tmp_path.glob("*.bak")).read_bytes()) == p


def test_atomic_failure_preserves_old(tmp_path, monkeypatch):
    import rf_link_calculator.persistence.json_io as module

    p = examples()["接收链路"]
    target = tmp_path / "p.json"
    save_project(p, target)
    before = target.read_bytes()

    def fail(*args):
        raise PermissionError("injected")

    monkeypatch.setattr(module.os, "replace", fail)
    with pytest.raises(PermissionError):
        save_project(replace(p, notes="new"), target)
    assert target.read_bytes() == before
    assert not list(tmp_path.glob(".rf-*"))


def test_csv_and_stable_ids():
    p = examples()["中文与特殊字符"]
    assert import_csv(export_csv(p), p) == p
    d = p.to_dict()
    ids = [s.id for s in p.stages]
    copied = edit_structure(d, "copy", ids[0])
    assert copied["stages"][1]["id"] not in ids
    moved = edit_structure(copied, "down", ids[0])
    assert moved["stages"][1]["id"] == ids[0]
    removed = edit_structure(moved, "delete", copied["stages"][1]["id"])
    assert [s["id"] for s in removed["stages"]] == ids
    assert safe_name("../CON") == "_CON" or "/" not in safe_name("../CON")


def test_excel_structure_and_injection():
    p = examples()["中文与特殊字符"]
    raw = build_excel(p, calculate(p))
    wb = load_workbook(BytesIO(raw))
    assert wb.sheetnames == list(SHEETS)
    assert wb["器件参数"]["D4"].value == p.stages[0].name
    assert wb["器件参数"]["D4"].data_type == "s"
    assert wb["器件参数"]["B203"].value is False
    assert wb["逐级计算"]["E203"].data_type == "f"
    assert "ISNUMBER" in wb["逐级计算"]["G4"].value
    assert "EXP(-ABS" in wb["逐级计算"]["T4"].value
    assert wb["压缩求解"].max_column == 205
    assert wb["压缩求解"]["C87"].data_type == "f"
    assert wb.calculation.fullCalcOnLoad
    assert wb["结果快照"]["A1"].value == "导出时快照，不随输入更新"
    for ws in wb:
        for row in ws:
            for cell in row:
                if cell.data_type == "f":
                    assert "#REF!" not in cell.value
    cached = load_workbook(BytesIO(raw), data_only=True)
    assert cached["系统汇总"]["B4"].value is None
    assert cached["结果快照"]["B6"].value is not None


def test_native_diagrams():
    p = examples()["下变频链路"]
    r = calculate(p)
    scene = build_scene(p, r, prefer_graphviz=False)
    svg = ET.fromstring(render_svg(scene))
    assert svg.tag.endswith("svg")
    assert not any(e.tag.endswith(("image", "script", "foreignObject")) for e in svg.iter())
    xml = ET.fromstring(render_drawio(scene))
    cells = xml.findall(".//mxCell")
    ids = {c.attrib["id"] for c in cells}
    edges = [c for c in cells if "source" in c.attrib]
    assert len(edges) == len(p.stages) + 2  # main chain plus LO
    for edge in edges:
        assert edge.attrib["source"] in ids and edge.attrib["target"] in ids
    assert any("triangle" in c.attrib.get("style", "") for c in cells)
    assert any(c.attrib["id"].startswith("lo-") for c in cells)
    assert not any("image=" in c.attrib.get("style", "") for c in cells)
    special = examples()["中文与特殊字符"]
    ET.fromstring(render_svg(build_scene(special, calculate(special))))
    ET.fromstring(render_drawio(build_scene(special, calculate(special))))


def test_segments_and_layout_failure(monkeypatch):
    p = examples()["200级长链路"]
    pages = segmented_scenes(p, calculate(p))
    assert len(pages) == 20
    assert len({n.stage_id for page in pages for n in page.nodes if n.stage_id}) == 200
    import rf_link_calculator.diagram.scene as module

    monkeypatch.setattr(module.shutil, "which", lambda x: "dot")

    def fail(*args, **kwargs):
        raise OSError("timeout simulation")

    monkeypatch.setattr(module.subprocess, "run", fail)
    scene = build_scene(p, calculate(p))
    assert scene.metadata["layout"] == "deterministic-linear"
    assert "失败" in scene.warnings[0]


def test_bundle_hash_and_failure(monkeypatch):
    p = examples()["接收链路"]
    r = calculate(p)
    with pytest.raises(ValueError):
        export_bundle(replace(p, notes="edited"), r)
    import rf_link_calculator.exporters.bundle as module

    def fail(*args):
        raise OSError("injected disk failure")

    monkeypatch.setattr(module, "build_excel", fail)
    raw, manifest = export_bundle(p, r)
    assert manifest["status"] == "partial_failure"
    with ZipFile(BytesIO(raw)) as z:
        assert "链路计算.xlsx" not in z.namelist()
        assert "Visio导入说明.txt" in z.namelist()
        m = json.loads(z.read("导出清单.json"))
        assert m["project_hash"] == r.project_hash
        assert not any(n.endswith("vsdx") for n in z.namelist())
