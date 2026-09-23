"""Advanced workbooks are explicit snapshots, never mislabeled scalar recalculators."""

import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def snapshot_workbook(project, result):
    wb = Workbook()
    info = wb.active
    info.title = "结果快照"
    info.append(["计算类型", "高级模型静态结果；修改工作簿不会重新计算"])
    info.append(["计算散列", result.calculation_hash])
    info.append(["公式版本", result.formula_version])
    info.append(["指标", "数值", "单位", "状态", "模型"])
    for k, m in result.metrics.items():
        info.append([k, m.value, m.unit, m.status, m.model])
    stages = wb.create_sheet("逐级结果")
    stages.append(["器件", "ID", "指标", "数值", "单位", "状态"])
    for row in result.stages:
        for k, m in row.metrics.items():
            stages.append([row.name, row.stage_id, k, m.value, m.unit, m.status])
    data = wb.create_sheet("模型数据")
    data.append(["器件ID", "模型", "来源", "数据SHA256", "行号", "数据"])
    for sid, slots in (project.lab or {}).get("models", {}).items():
        for slot, ds in slots.items():
            meta = ds.get("metadata", {})
            arrays = {k: v for k, v in ds.items() if isinstance(v, list)}
            for i in range(max([1] + [len(v) for v in arrays.values()])):
                row = {k: v[i] for k, v in arrays.items() if i < len(v)}
                data.append(
                    [
                        sid,
                        slot + ":" + ds["kind"],
                        str(meta.get("filename", "")),
                        meta.get("sha256", meta.get("dataset_sha256", "")),
                        i + 1,
                        json.dumps(row, ensure_ascii=False),
                    ]
                )
    notes = wb.create_sheet("假设与范围")
    for text in result.assumptions:
        notes.append([text])
    for issue in result.issues:
        notes.append([issue.code, issue.message])
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="087F8C")
        for col in "ABCDEF":
            sheet.column_dimensions[col].width = 28 if col != "F" else 80
        # Imported names and file metadata are text, never spreadsheet formulas.
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
