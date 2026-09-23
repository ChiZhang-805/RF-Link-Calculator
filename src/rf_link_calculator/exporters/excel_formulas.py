"""An independent Excel implementation. Empty means unavailable, never implicit zero."""

from openpyxl.utils import get_column_letter

FIRST, LAST = 4, 203


def guard(expression: str, *dependencies: str) -> str:
    checks = ",".join(f"ISNUMBER({x})" for x in dependencies)
    return f'=IF(AND({checks}),{expression},"")'


def loss(x: str, ip1: str, p: str) -> str:
    z = f"(LN(10^({p}/10)-1)+{p}*({x}-{ip1})*LN(10)/10)"
    return f"(10/({p}*LN(10))*(MAX({z},0)+LN(1+EXP(-ABS({z})))))"


def row_formulas(r: int) -> dict[str, str]:
    q = r - 1
    s = lambda col: f"'器件参数'!{col}{r}"
    enabled, gain, mode = s("B"), s("F"), s("G")
    valid_gain = f"AND(ISNUMBER({gain}),ABS({gain})<=300)"
    passive_types = ",".join(
        f'{s("E")}="{k}"'
        for k in (
            "filter_bpf",
            "filter_lpf",
            "filter_hpf",
            "filter_bsf",
            "attenuator",
            "cable",
            "connector",
            "switch",
            "isolator",
            "coupler",
            "splitter",
        )
    )
    mixer_ok = f'OR({s("E")}<>"mixer",AND(OR({s("V")}="SSB",{s("V")}="DSB"),{s("W")}=TRUE))'
    noise = (
        f'IF(NOT({enabled}),1,IF(NOT({mixer_ok}),"",IF({mode}="ideal",1,'
        f'IF(AND({mode}="manual",ISNUMBER({s("H")}),{s("H")}>=0),10^({s("H")}/10),'
        f'IF(AND({mode}="passive_thermal",OR({passive_types}),{valid_gain},{gain}<=0,ISNUMBER({s("I")}),{s("I")}>0),'
        f'1+(10^(-{gain}/10)-1)*{s("I")}/Tref_K,"")))))'
    )
    ip = f'IF(AND({s("J")}="finite",ISNUMBER({s("L")}),ISNUMBER(C{r})),IF({s("K")}="input",{s("L")},IF({s("K")}="output",{s("L")}-C{r},"")),"")'
    ip1 = f'IF(AND({s("M")}="finite",ISNUMBER({s("O")}),ISNUMBER(C{r})),IF({s("N")}="input",{s("O")},IF({s("N")}="actual_output",{s("O")}-C{r}+1,IF({s("N")}="linear_output",{s("O")}-C{r},""))),"")'
    f = {
        "A": f"={s('A')}",
        "B": f"={s('D')}",
        "C": f'=IF(NOT({enabled}),0,IF({valid_gain},{gain},""))',
        "D": f'=IF(ISNUMBER(E{q}),E{q},"")',
        "E": guard(f"C{r}+D{r}", f"C{r}", f"D{r}"),
        "F": guard(f"10^(D{r}/10)", f"D{r}"),
        "G": f"={noise}",
        "H": guard(f"H{q}+(G{r}-1)/F{r}", f"H{q}", f"G{r}", f"F{r}"),
        "I": guard(f"10*LOG10(H{r})", f"H{r}"),
        "J": guard(f"Pin_dBm+D{r}", "Pin_dBm", f"D{r}"),
        "K": guard(f"Pin_dBm+E{r}", "Pin_dBm", f"E{r}"),
        "L": f"={ip}",
        "M": f'=IF(NOT(ISNUMBER(M{q})),"",IF(OR(NOT({enabled}),{s("J")}="ideal"),M{q},IF(AND(ISNUMBER(L{r}),ISNUMBER(F{r})),M{q}+F{r}/10^(L{r}/10),"")))',
        "N": f'=IF(AND(ISNUMBER(M{r}),M{r}>0),-10*LOG10(M{r}),"")',
        "O": guard(f"N{r}+E{r}", f"N{r}", f"E{r}"),
        "P": f"={ip1}",
        "Q": f'=IF({s("P")}="",IF(AND(ISNUMBER(DefaultP),DefaultP>=1,DefaultP<=10),DefaultP,""),IF(AND(ISNUMBER({s("P")}),{s("P")}>=1,{s("P")}<=10),{s("P")},""))',
        "R": f'=IF(ISNUMBER(U{q}),U{q},"")',
        "S": guard(f"LN(10^(Q{r}/10)-1)+Q{r}*(R{r}-P{r})*LN(10)/10", f"Q{r}", f"R{r}", f"P{r}"),
        "T": f'=IF(X{r}=0,0,IF(AND(X{r}=1,ISNUMBER(S{r})),10/(Q{r}*LN(10))*(MAX(S{r},0)+LN(1+EXP(-ABS(S{r})))),""))',
        "U": guard(f"R{r}+C{r}-T{r}", f"R{r}", f"C{r}", f"T{r}"),
        "V": guard(f"Pin_dBm+E{r}-U{r}", "Pin_dBm", f"E{r}", f"U{r}"),
        "W": f'=IF(AND({enabled},X{r}=1,ISNUMBER(D{r})),P{r}-D{r},"")',
        "X": f'=IF(NOT(ISNUMBER(C{r})),-1,IF(OR(NOT({enabled}),{s("M")}="ideal"),0,IF(AND(ISNUMBER(P{r}),ISNUMBER(Q{r})),1,-1)))',
        "Y": f'=IF(ISNUMBER(H{r}),"有效","未知／噪声参数或口径不完整")',
        "Z": f'=IF(NOT(ISNUMBER(M{r})),"未知／IP3参数不完整",IF(M{r}=0,"理想／无有限值","有效"))',
        "AA": f'=IF(ISNUMBER(U{r}),"模型估算","未知／增益或压缩参数不完整")',
        "AB": f'=IF(ISNUMBER(AC{q}),AC{q},"")',
        "AC": f'=IF(NOT(ISNUMBER(AB{r})),"",IF(OR(NOT({enabled}),{s("E")}<>"mixer"),AB{r},IF(AND(ISNUMBER({s("T")}),{s("T")}>0),IF({s("U")}="sum",AB{r}+{s("T")},IF({s("U")}="difference",ABS(AB{r}-{s("T")}),"")),"")))',
        "AD": guard(f"10*LOG10(G{r})", f"G{r}"),
        "AE": guard(f"Tref_K*(G{r}-1)/F{r}", f"G{r}", f"F{r}"),
        "AF": f'=IF(AND(ISNUMBER(H{r}),ISNUMBER(E{r}),NoiseConfigOK),10*LOG10(Boltzmann*Bnoise_Hz*(Tsource_K+Tref_K*(H{r}-1)))+30+E{r},"")',
        "AG": guard(f"K{r}-AF{r}", f"K{r}", f"AF{r}"),
        "AH": f'=IF(AND(ToneEnabled,ISNUMBER(E{r}),ISNUMBER(TonePin)),TonePin+E{r},"")',
        "AI": f'=IF(AND(ToneEnabled,ISNUMBER(E{r}),ISNUMBER(N{r}),ISNUMBER(TonePin)),3*TonePin+E{r}-2*N{r},"")',
        "AJ": f'=IF(X{r}=1,Q{r},"")',
        "AK": f'=IF(NOT({enabled}),"旁路",IF(NOT(ISNUMBER(AC{r})),"频率未知",IF(AC{r}<=SignalBW/2,"跨零频",IF(AND(ISNUMBER({s("R")}),ISNUMBER({s("S")})),IF(AND(AB{r}-SignalBW/2>={s("R")},AB{r}+SignalBW/2<={s("S")}),"有效","超声明通带"),"未声明通带"))))',
        "AL": f"=IF(AND(X{r}=1,ISNUMBER(D{r})),10^(Q{r}*(D{r}-P{r})/10),0)",
        "AM": f'=IF(AND({enabled},ISNUMBER({s("X")}),OR(AND(ISNUMBER(J{r}),J{r}>{s("X")}),IF(AND(ToneEnabled,ISNUMBER(TonePin),ISNUMBER(D{r})),TonePin+10*LOG10(2)+D{r}>{s("X")},FALSE))),"超过绝对最大输入额定值","")',
    }
    # Track both fundamentals and both IM3 frequencies through each selected mixer.
    for col in ("AN", "AO", "AP", "AQ"):
        prev = f"{col}{q}"
        f[col] = (
            f'=IF(NOT(ISNUMBER({prev})),"",IF(OR(NOT({enabled}),{s("E")}<>"mixer"),{prev},IF(AND(ISNUMBER({s("T")}),{s("T")}>0),IF({s("U")}="sum",{prev}+{s("T")},ABS({prev}-{s("T")})),"")))'
        )
    f["AR"] = (
        f"=IF(NOT(AR{q}),FALSE,IF(NOT({enabled}),TRUE,AND(COUNT(AN{r}:AQ{r})=4,MIN(AN{r}:AQ{r})>0,"
        f"IF(AND(ISNUMBER({s('R')}),ISNUMBER({s('S')})),AND(MIN(AN{q}:AQ{q})>={s('R')},MAX(AN{q}:AQ{q})<={s('S')}),TRUE),"
        f'IF(AND({s("E")}="mixer",{s("U")}="difference"),OR({s("T")}<MIN(AN{q}:AQ{q}),{s("T")}>MAX(AN{q}:AQ{q})),TRUE))))'
    )
    # Expected numeric-domain failures are made visible; no false valid zero.
    for col in ("F", "G", "H", "M", "AL"):
        f[col] = f'=IFERROR({f[col][1:]},"数值域错误")'
    return f


def add_solver(ws):
    ws["A1"] = "系统1 dB压缩：区间检查＋80次无宏二分；全部理想／缺参数分别标记"
    ws["A2"] = "完整模型状态"
    ws["B2"] = (
        '=IF(COUNTIF(\'逐级计算\'!X4:X203,-1)>0,"未知",IF(COUNTIF(\'逐级计算\'!X4:X203,1)=0,"理想","可求解"))'
    )
    ws["A3"] = "同p闭式对照"
    ws["B3"] = (
        "=IF(AND(B2=\"可求解\",MIN('逐级计算'!AJ4:AJ203)=MAX('逐级计算'!AJ4:AJ203),SUM('逐级计算'!AL4:AL203)>0),-10/MIN('逐级计算'!AJ4:AJ203)*LOG10(SUM('逐级计算'!AL4:AL203)),\"\")"
    )
    ws["A4"] = "夹根状态"
    ws["B4"] = (
        '=IF(B2<>"可求解",B2,IF(AND(ISNUMBER(D6),ISNUMBER(D7),D6<1,D7>1),"已夹根","未夹根／数值域错误"))'
    )
    for col, label in enumerate(("步骤", "下界", "试探输入dBm", "总压缩dB"), 1):
        ws.cell(5, col, label)
    # Follows the full chain, including all 200 allocated input rows.
    lastcol = get_column_letter(204)
    uppercol = get_column_letter(205)
    ws.cell(5, 205, "上界")
    for r in range(6, 88):
        ws.cell(r, 1, "下界验证" if r == 6 else ("上界验证" if r == 7 else r - 7))
        if r == 6:
            ws.cell(r, 2, '=IF($B$2="可求解",MIN(\'逐级计算\'!W4:W203)-80,"")')
            ws.cell(r, 205, '=IF($B$2="可求解",MAX(\'逐级计算\'!W4:W203)+20,"")')
            ws.cell(r, 3, '=IF(ISNUMBER(B6),B6,"")')
        elif r == 7:
            ws.cell(r, 3, f'=IF(ISNUMBER({uppercol}6),{uppercol}6,"")')
        elif r == 8:
            ws.cell(r, 2, '=IF($B$4="已夹根",B6,"")')
            ws.cell(r, 205, f'=IF($B$4="已夹根",{uppercol}6,"")')
            ws.cell(r, 3, guard(f"(B{r}+{uppercol}{r})/2", f"B{r}", f"{uppercol}{r}"))
        else:
            ws.cell(r, 2, f'=IF(ISNUMBER(D{r - 1}),IF(D{r - 1}>1,B{r - 1},C{r - 1}),"")')
            ws.cell(r, 205, f'=IF(ISNUMBER(D{r - 1}),IF(D{r - 1}>1,C{r - 1},{uppercol}{r - 1}),"")')
            ws.cell(r, 3, guard(f"(B{r}+{uppercol}{r})/2", f"B{r}", f"{uppercol}{r}"))
        prev = f"C{r}"
        for i in range(200):
            sr, c = i + 4, i + 5
            base = lambda col: f"'逐级计算'!${col}${sr}"
            expression = f'IF(AND(ISNUMBER({prev}),{base("X")}>-1),IF({base("X")}=0,{prev}+{base("C")},{prev}+{base("C")}-{loss(prev, base("P"), base("Q"))}),"")'
            ws.cell(r, c, "=" + expression)
            if r == 6:
                ws.cell(5, c, f"第{i + 1}级输出")
            prev = f"{get_column_letter(c)}{r}"
        ws.cell(
            r, 4, guard(f"C{r}+'逐级计算'!$E$203-{lastcol}{r}", f"C{r}", "'逐级计算'!$E$203", f"{lastcol}{r}")
        )
    ws["A89"] = "二分与闭式差异dB"
    ws["B89"] = guard("ABS(C87-B3)", "C87", "B3")
    ws["A90"] = "数值一致性"
    ws["B90"] = '=IF(ISNUMBER(B89),IF(B89<0.00001,"通过","公式偏差，禁止采信"),"不同p或无可用闭式")'
