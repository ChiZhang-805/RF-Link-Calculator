"""Editable 200-stage workbook plus an explicitly historical Python snapshot."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from rf_link_calculator.domain.labels import METRICS, STATUS
from rf_link_calculator.domain.models import Project, ResultBundle
from rf_link_calculator.persistence.json_io import hashes

from .excel_formulas import FIRST, LAST, add_solver, guard, row_formulas

SHEETS = ("使用说明", "输入配置", "器件参数", "逐级计算", "系统汇总", "压缩求解", "结果快照", "公式与来源")
INPUT_HEADERS = [
    "顺序",
    "启用",
    "器件ID",
    "名称",
    "类型",
    "增益dB",
    "噪声模式",
    "NF dB",
    "物理温度K",
    "IP3模式",
    "IP3参考",
    "IP3 dBm",
    "P1模式",
    "P1参考",
    "P1 dBm",
    "形状p",
    "来源与条件",
    "频率下限Hz",
    "频率上限Hz",
    "LO Hz",
    "变频关系",
    "NF口径",
    "噪声口径已确认",
    "最大输入dBm",
    "备注",
]
CALC_HEADERS = [
    "顺序",
    "器件名称",
    "有效增益dB",
    "前级增益dB",
    "累计增益dB",
    "前级线性增益",
    "本级噪声因子",
    "累计噪声因子",
    "累计NF dB",
    "线性输入dBm",
    "线性输出dBm",
    "本级IIP3 dBm",
    "IP3倒数和",
    "累计IIP3 dBm",
    "累计OIP3 dBm",
    "本级输入P1 dBm",
    "有效p",
    "估算输入dBm",
    "稳定指数z",
    "本级压缩dB",
    "估算输出dBm",
    "累计压缩dB",
    "入口折算阈值dBm",
    "压缩模式 -1未知/0理想/1有限",
    "噪声状态",
    "IP3状态",
    "压缩状态",
    "输入频率Hz",
    "输出频率Hz",
    "本级NF dB",
    "输入温度贡献K",
    "输出噪声dBm",
    "SNR dB",
    "每音输出dBm",
    "弱非线性IM3外推dBm",
    "有限级p",
    "频率适用性",
    "同p闭式项",
    "额定告警",
    "基波1 Hz",
    "基波2 Hz",
    "IM3低侧Hz",
    "IM3高侧Hz",
    "双音通带条件",
]


def plain(cell, value):
    """Untrusted text is explicitly a string, including text starting with =, +, -, @."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def build_excel(project: Project, result: ResultBundle) -> bytes:
    if hashes(project)[0] != result.project_hash:
        raise ValueError("输入与结果快照不一致，拒绝导出过期结果")
    wb = Workbook()
    wb.remove(wb.active)
    for title in SHEETS:
        wb.create_sheet(title)
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    info, conf, inputs, calc, summary, solver, snapshot, formula = [wb[x] for x in SHEETS]
    info_rows = [
        "RF-Link-Calculator｜可重算射频链路工作簿",
        project.project_name + " / " + project.link_name,
        "浅蓝色为可编辑输入，浅灰色为真实公式。打开Excel后请完整重算（Ctrl+Alt+F9）。",
        "openpyxl不执行公式；本文件未在本机Excel桌面重算验证。结果快照是导出时静态报告。",
        "最多200级，第4～203行。取消隐藏空行后填写并启用；删除器件请禁用，勿删除物理行。",
        "顺序由器件参数表的整行位置决定；整表排序/移动需包含所有列。不得单独排序某参数列。",
        "噪声模式: manual/passive_thermal/ideal/unknown；IP3与P1: finite/ideal/unknown。",
        "IP3参考: input/output；P1参考: input/actual_output/linear_output。实际输出P1换算有额外1 dB。",
        "不启用任何宏、插件或循环计算。压缩求解表使用82次完整链路传播（2次夹根检查＋80次二分）。",
        "JSON是标准回读格式；Excel修改不会回写网页、JSON或导出时快照。",
        "NF、IP3、P1未知分别阻断依赖指标；双音每音功率独立于主CW输入。",
        "Excel线性因子有数值范围，极端链路会显示数值域错误；Python采用对数域可覆盖更大范围。",
        "频率适用性和额定告警参见逐级计算AK/AM列；IP3/IM3为独立弱非线性模型。",
        "禁用代表逻辑旁路；边界标注不能参与天线增益或ADC量化噪声预算。",
        f"输入散列: {result.project_hash}",
        f"计算散列: {result.calculation_hash}",
        f"公式版本: {result.formula_version}；模型: {result.compression_model_version}",
    ]
    for r, value in enumerate(info_rows, 1):
        plain(info.cell(r, 1), value)
    a, t = project.analysis, project.analysis.two_tone
    cfg = [
        ("Pin_dBm", "主CW输入功率", a.input_power_dbm, "dBm"),
        ("SourceFreq", "输入中心频率", a.source_frequency_hz, "Hz"),
        ("SignalBW", "信号带宽", a.signal_bandwidth_hz, "Hz"),
        ("Bnoise_Hz", "等效噪声带宽", a.noise_bandwidth_hz, "Hz"),
        ("Tsource_K", "输入源等效噪声温度", a.source_noise_temperature_k, "K"),
        ("Tref_K", "固定NF参考温度（勿改）", 290, "K"),
        ("RequiredSNR_dB", "所需信噪比", a.required_snr_db, "dB"),
        ("ImplLoss_dB", "实现损失", a.implementation_loss_db, "dB"),
        ("Backoff_dB", "压缩回退", a.compression_backoff_db, "dB"),
        ("DefaultP", "默认压缩形状p", a.default_compression_p, ""),
        ("ToneEnabled", "独立双音分析启用", t.enabled, "bool"),
        ("TonePin", "每音输入功率", t.each_tone_power_dbm, "dBm"),
        ("ToneSpacing", "双音间隔", t.spacing_hz, "Hz"),
        ("Boltzmann", "玻尔兹曼常数（勿改）", 1.380649e-23, "J/K"),
    ]
    conf.append(["内部名称", "中文说明", "输入值", "单位"])
    for r, (key, label, value, unit) in enumerate(cfg, 2):
        for c, v in enumerate((key, label, value, unit), 1):
            plain(conf.cell(r, c), v)
        wb.defined_names.add(DefinedName(key, attr_text=f"'输入配置'!$C${r}"))
    conf["A17"], conf["B17"] = "NoiseConfigOK", "噪声配置有效"
    conf["C17"] = (
        "=AND(ISNUMBER(Bnoise_Hz),Bnoise_Hz>0,ISNUMBER(Tsource_K),Tsource_K>0,Tref_K=290,Boltzmann=1.380649E-23)"
    )
    wb.defined_names.add(DefinedName("NoiseConfigOK", attr_text="'输入配置'!$C$17"))
    for refs, lower, upper in (("C3:C6 C14", 1e-12, 1e15), ("C11", 1, 10)):
        dv = DataValidation(type="decimal", operator="between", formula1=lower, formula2=upper)
        dv.errorTitle, dv.error, dv.showErrorMessage = (
            "数值范围错误",
            f"请输入 {lower}～{upper} 范围内的数值",
            True,
        )
        conf.add_data_validation(dv)
        for ref in refs.split():
            dv.add(ref)
    inputs["A1"] = "编辑输入：最多200级；空行明确禁用，整行重排；名称与备注不会执行为公式"
    calc["A1"] = "当前Excel公式结果（自动重算）｜压缩为模型估算｜噪声为小信号预算"
    for c, header in enumerate(INPUT_HEADERS, 1):
        inputs.cell(3, c, header)
    for c, header in enumerate(CALC_HEADERS, 1):
        calc.cell(2, c, header)
    calc["E3"], calc["H3"], calc["M3"], calc["U3"], calc["AC3"] = 0, 1, 0, "=Pin_dBm", "=SourceFreq"
    for col, delta in (("AN", -0.5), ("AO", 0.5), ("AP", -1.5), ("AQ", 1.5)):
        calc[f"{col}3"] = f"=SourceFreq+({delta})*ToneSpacing"
    calc["AR3"] = "=AND(MIN(AN3:AQ3)>0,ToneSpacing>0)"
    for r in range(FIRST, LAST + 1):
        index = r - FIRST
        if index < len(project.stages):
            s = project.stages[index]
            values = [
                s.order,
                s.enabled,
                s.id,
                s.name,
                s.type,
                s.gain_db,
                s.noise.mode,
                s.noise.value_db,
                s.physical_temperature_k,
                s.ip3.mode,
                s.ip3.reference,
                s.ip3.value_dbm,
                s.p1db.mode,
                s.p1db.reference,
                s.p1db.value_dbm,
                s.compression_p,
                str(s.source),
                *(s.frequency_range_hz or (None, None)),
                s.mixer.lo_frequency_hz if s.mixer else None,
                s.mixer.relation if s.mixer else "",
                s.mixer.noise_convention if s.mixer else "",
                s.mixer.noise_compatible if s.mixer else False,
                s.absolute_max_input_dbm,
                s.notes,
            ]
        else:
            values = [
                index + 1,
                False,
                f"reserved-{index + 1}",
                "",
                "custom",
                None,
                "unknown",
                None,
                290,
                "unknown",
                "input",
                None,
                "unknown",
                "input",
                None,
                None,
                "",
                None,
                None,
                None,
                "",
                "unknown",
                False,
                None,
                "",
            ]
        for c, v in enumerate(values, 1):
            plain(inputs.cell(r, c), v)
        for col, expr in row_formulas(r).items():
            calc[f"{col}{r}"] = expr
        if index >= len(project.stages):
            inputs.row_dimensions[r].hidden = True
            calc.row_dimensions[r].hidden = True
    table = Table(displayName="StageInputs", ref=f"A3:Y{LAST}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    inputs.add_table(table)
    for col, options in {
        "G": "unknown,manual,passive_thermal,ideal",
        "J": "unknown,finite,ideal",
        "M": "unknown,finite,ideal",
        "K": "input,output",
        "N": "input,actual_output,linear_output",
        "U": "sum,difference",
        "V": "unknown,SSB,DSB",
    }.items():
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=False)
        dv.errorTitle, dv.error, dv.showErrorMessage = "参数模式错误", "请选择列表中的值", True
        inputs.add_data_validation(dv)
        dv.add(f"{col}{FIRST}:{col}{LAST}")
    add_solver(solver)
    summary["A1"] = "当前Excel公式结果｜P1与压缩为模型估算｜空值必须结合状态列阅读"
    summary.append(["指标", "当前值", "单位", "状态／条件", "内部键"])
    rows = {key: i + 4 for i, key in enumerate(METRICS)}
    cell = lambda key: f"B{rows[key]}"
    state = lambda key: f"D{rows[key]}"
    tail = lambda col: f"'逐级计算'!{col}{LAST}"
    formulas = {}
    for key, col in {
        "gain_db": "E",
        "linear_output_dbm": "K",
        "nf_db": "I",
        "noise_output_dbm": "AF",
        "snr_db": "AG",
        "iip3_dbm": "N",
        "oip3_dbm": "O",
        "compressed_output_dbm": "U",
        "compression_db": "V",
        "frequency_output_hz": "AC",
        "tone_output_dbm": "AH",
    }.items():
        formulas[key] = f'=IF(ISNUMBER({tail(col)}),{tail(col)},"")'
    formulas["te_k"] = guard(f"Tref_K*({tail('H')}-1)", tail("H"))
    formulas["noise_input_dbm"] = guard(
        f"{cell('noise_output_dbm')}-{cell('gain_db')}", cell("noise_output_dbm"), cell("gain_db")
    )
    formulas["ip1_dbm"] = "=IF('压缩求解'!B4=\"已夹根\",'压缩求解'!C87,\"\")"
    formulas["op1_dbm"] = guard(f"{cell('ip1_dbm')}+{cell('gain_db')}-1", cell("ip1_dbm"), cell("gain_db"))
    formulas["sensitivity_dbm"] = guard(
        f"{cell('noise_input_dbm')}+RequiredSNR_dB+ImplLoss_dB",
        cell("noise_input_dbm"),
        "RequiredSNR_dB",
        "ImplLoss_dB",
    )
    formulas["compression_dr_db"] = guard(
        f"{cell('ip1_dbm')}-Backoff_dB-{cell('sensitivity_dbm')}",
        cell("ip1_dbm"),
        "Backoff_dB",
        cell("sensitivity_dbm"),
    )
    formulas["sfdr3_db"] = guard(
        f"2/3*({cell('iip3_dbm')}-{cell('noise_input_dbm')})", cell("iip3_dbm"), cell("noise_input_dbm")
    )
    formulas["sfdr3_normalized"] = guard(f"{cell('sfdr3_db')}+20/3*LOG10(Bnoise_Hz)", cell("sfdr3_db"))
    formulas["tone_limit_dbm"] = guard(
        f"({cell('noise_input_dbm')}+2*{cell('iip3_dbm')})/3", cell("noise_input_dbm"), cell("iip3_dbm")
    )
    formulas["tone_span_db"] = guard(
        f"{cell('tone_limit_dbm')}-{cell('sensitivity_dbm')}", cell("tone_limit_dbm"), cell("sensitivity_dbm")
    )
    formulas["tone_input_dbm"] = '=IF(AND(ToneEnabled,ISNUMBER(TonePin)),TonePin,"")'
    formulas["tone_total_dbm"] = '=IF(AND(ToneEnabled,ISNUMBER(TonePin)),TonePin+10*LOG10(2),"")'
    tone_gate = f"AND(ToneEnabled,ISNUMBER(TonePin),{tail('AR')},ISNUMBER({cell('iip3_dbm')}),ISNUMBER({cell('gain_db')}),IF(ISNUMBER({cell('ip1_dbm')}),IF(AND(ToneEnabled,ISNUMBER(TonePin)),TonePin+10*LOG10(2)<{cell('ip1_dbm')}-Backoff_dB,TRUE),TRUE))"
    formulas["im3_dbm"] = f'=IF({tone_gate},3*TonePin+{cell("gain_db")}-2*{cell("iip3_dbm")},"")'
    formulas["im3_dbc"] = f'=IF({tone_gate},-2*({cell("iip3_dbm")}-TonePin),"")'
    for key, r in rows.items():
        label, unit = METRICS[key]
        summary.cell(r, 1, label)
        summary.cell(r, 2, formulas[key])
        summary.cell(r, 3, unit)
        summary.cell(r, 5, key)
        msg = (
            "模型估算"
            if key in ("ip1_dbm", "op1_dbm", "compressed_output_dbm", "compression_db", "compression_dr_db")
            else "有效"
        )
        summary.cell(r, 4, f'=IF(ISNUMBER(B{r}),"{msg}","未知／依赖不完整或数值域错误")')
    for key in ("iip3_dbm", "oip3_dbm"):
        summary[state(key)] = f"={tail('Z')}"
    for key in ("ip1_dbm", "op1_dbm"):
        summary[state(key)] = "=IF('压缩求解'!B4=\"已夹根\",\"模型估算\",'压缩求解'!B4)"
    for key in ("im3_dbm", "im3_dbc"):
        summary[state(key)] = (
            f'=IF(NOT(ToneEnabled),"未启用",IF(ISNUMBER({cell(key)}),IF(\'压缩求解\'!B2="未知","弱非线性外推；压缩有效区未核实","弱非线性外推"),"不适用／通带、压缩或IP3条件不满足"))'
        )
    summary["A36"] = "负动态范围／额定告警"
    summary["B36"] = (
        f'=IF(OR({cell("compression_dr_db")}<0,{cell("sfdr3_db")}<0,COUNTIF(\'逐级计算\'!AM4:AM203,"超过绝对最大输入额定值")>0),"检查动态范围与额定值","")'
    )
    snapshot.append(["导出时快照，不随输入更新"])
    snapshot.append(["导出时间", result.created_at])
    snapshot.append(["输入散列", result.project_hash])
    snapshot.append(["计算散列", result.calculation_hash])
    snapshot.append(["指标", "Python导出值", "单位", "状态", "原因"])
    for key, v in result.metrics.items():
        snapshot.append([METRICS.get(key, (key, ""))[0], v.value, v.unit, STATUS[v.status], v.reason])
    snapshot.append(["逐级快照（静态）"])
    snapshot.append(["器件", "线性输入dBm", "线性输出dBm", "累计NF dB", "累计IIP3 dBm", "压缩输出估算dBm"])
    for s in result.stages:
        row = [s.name] + [
            s.metrics[k].value
            for k in ("linear_input_dbm", "linear_output_dbm", "nf_db", "iip3_dbm", "compressed_output_dbm")
        ]
        snapshot.append(row)
        plain(snapshot.cell(snapshot.max_row, 1), s.name)
    formula_rows = [
        "公式、参考点与模型",
        "Fsys=1+Σ(Fi-1)/A_before；NF=10log10(Fsys)",
        "Te=290(Fsys-1)；Nin=k*Bn*(Ts+Te)；Nout=Nin*Gsys",
        "1/IIP3sys=ΣA_before/IIP3_i（mW）",
        "IM3out=3*每音Pin+Gsys-2*IIP3sys；SFDR3=2/3*(IIP3sys-Nin)",
        "实际OP1=IP1+G-1；线性外推OP1=IP1+G；C(x)=10/p log10[1+(10^(p/10)-1)10^(p(x-IP1)/10)]",
        "系统P1求解总压缩1 dB；同p闭式只作为独立对照，不将最弱阈值冒充系统P1。",
        "https://openpyxl.readthedocs.io/en/stable/simple_formulae.html",
        "https://www.mathworks.com/help/rf/ref/rfbudget.html",
        "https://www.analog.com/en/resources/technical-articles/use-selectivity-to-improve-receiver-intercept-point.html",
        *result.assumptions,
        *[f"{i.code} {i.stage_id or ''} {i.field_path}: {i.message}" for i in result.issues],
    ]
    for r, v in enumerate(formula_rows, 1):
        plain(formula.cell(r, 1), v)
    # Compact styling avoids formatting tens of thousands of solver cells individually.
    for ws in wb:
        ws.freeze_panes = "C4" if ws in (inputs, calc) else "B4"
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", ws.PAPERSIZE_A3
        ws.print_title_rows = "1:3"
        for c in range(1, min(ws.max_column, 45) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 20
        ws.column_dimensions["A"].width = 40
        for cell_ in ws[1]:
            cell_.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=13)
            cell_.fill = PatternFill("solid", fgColor="164E63")
        if ws != solver:
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(vertical="top", wrap_text=True)
                    c.number_format = (
                        "0.000" if isinstance(c.value, (int, float)) or c.data_type == "f" else "General"
                    )
                    if ws in (inputs, conf):
                        c.fill = PatternFill("solid", fgColor="E0F2FE")
                    elif c.data_type == "f":
                        c.fill = PatternFill("solid", fgColor="F1F5F9")
        if ws in (info, formula):
            ws.column_dimensions["A"].width = 120
        if ws in (summary, snapshot):
            ws.column_dimensions["B"].width = 35
        if ws == summary:
            for key, r in rows.items():
                unit = METRICS[key][1]
                ws.cell(r, 2).number_format = f'0.000 "{unit}"'
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
